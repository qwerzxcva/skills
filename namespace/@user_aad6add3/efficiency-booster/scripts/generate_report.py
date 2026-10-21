#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
效率提升器 - Word报告生成器
根据四步框架生成效率提升方案和执行手册
"""

import sys
import json
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    print("[ERROR] Please install: pip install python-docx")
    sys.exit(1)


def create_cover(doc, task_name, date_str):
    """创建封面"""
    # 空行
    for _ in range(6):
        doc.add_paragraph()
    
    # 主标题
    title_para = doc.add_paragraph()
    title_run = title_para.add_run("效率提升方案")
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 128, 128)  # 青色
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 空行
    for _ in range(2):
        doc.add_paragraph()
    
    # 副标题
    subtitle_para = doc.add_paragraph()
    subtitle_run = subtitle_para.add_run("目标-路径-资源-复盘 四步闭环")
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.color.rgb = RGBColor(102, 102, 102)
    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 任务名称
    if task_name:
        task_para = doc.add_paragraph()
        task_run = task_para.add_run(f"任务：{task_name}")
        task_run.font.size = Pt(14)
        task_run.font.color.rgb = RGBColor(51, 51, 51)
        task_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 空行
    for _ in range(4):
        doc.add_paragraph()
    
    # 日期
    date_para = doc.add_paragraph()
    date_run = date_para.add_run(date_str)
    date_run.font.size = Pt(12)
    date_run.font.color.rgb = RGBColor(128, 128, 128)
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 分页
    doc.add_page_break()


def add_heading_custom(doc, text, level=1):
    """添加自定义标题"""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        if level == 1:
            run.font.color.rgb = RGBColor(0, 128, 128)  # 青色
            run.font.size = Pt(18)
        elif level == 2:
            run.font.color.rgb = RGBColor(0, 153, 153)  # 深青色
            run.font.size = Pt(14)
        else:
            run.font.color.rgb = RGBColor(51, 51, 51)
    return heading


def add_step_section(doc, step_num, step_name, step_key, content, color):
    """添加步骤章节"""
    # 步骤标题
    heading = add_heading_custom(doc, f"第{step_num}步：{step_name}", level=2)
    
    # 内容
    if content:
        for key, value in content.items():
            if value:
                # 子标题
                para = doc.add_paragraph()
                label = para.add_run(f"{key}：")
                label.font.bold = True
                label.font.color.rgb = RGBColor(int(color[1:3], 16), int(color[3:5], 16), int(color[5:7], 16))
                
                # 内容
                if isinstance(value, list):
                    for item in value:
                        p = doc.add_paragraph(style='List Bullet')
                        p.add_run(str(item))
                        p.paragraph_format.line_spacing = 1.5
                else:
                    content_para = doc.add_paragraph(str(value))
                    content_para.paragraph_format.line_spacing = 1.5
    
    doc.add_paragraph()


def create_5w1h_table(doc, clarify_data):
    """创建5W1H分析表"""
    add_heading_custom(doc, "5W1H分析", level=2)
    
    # 创建表格
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Grid Accent 1'
    
    # 表头
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '维度'
    hdr_cells[1].text = '内容'
    
    # 设置表头样式
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 添加数据行
    w5h1_items = [
        ("Why (目的)", clarify_data.get('why', '')),
        ("What (标准)", clarify_data.get('what', '')),
        ("Who (责任)", clarify_data.get('who', '')),
        ("When (节点)", clarify_data.get('when', '')),
        ("Where (边界)", clarify_data.get('where', '')),
        ("How (方法)", clarify_data.get('how', '')),
    ]
    
    for label, value in w5h1_items:
        if value:
            row_cells = table.add_row().cells
            row_cells[0].text = label
            row_cells[1].text = str(value)
    
    doc.add_paragraph()


def create_wbs_section(doc, deconstruct_data):
    """创建WBS分解"""
    add_heading_custom(doc, "WBS任务分解", level=2)
    
    # 阶段列表
    phases = deconstruct_data.get('phases', [])
    if phases:
        for phase in phases:
            para = doc.add_paragraph()
            label = para.add_run(f"阶段：{phase.get('name', '')}")
            label.font.bold = True
            label.font.size = Pt(12)
            
            tasks = phase.get('tasks', [])
            for task in tasks:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(task)
                p.paragraph_format.line_spacing = 1.5
    
    doc.add_paragraph()


def create_optimization_section(doc, optimize_data):
    """创建优化方案"""
    add_heading_custom(doc, "路径优化", level=2)
    
    # 二八分析
    pareto = optimize_data.get('pareto', {})
    if pareto:
        para = doc.add_paragraph()
        label = para.add_run("二八分析：")
        label.font.bold = True
        
        core_tasks = pareto.get('core_tasks', [])
        if core_tasks:
            p = doc.add_paragraph("20%核心任务（高影响力）：")
            for task in core_tasks:
                bullet = doc.add_paragraph(style='List Bullet')
                bullet.add_run(task)
    
    # 假动作剔除
    fake_actions = optimize_data.get('fake_actions', [])
    if fake_actions:
        para = doc.add_paragraph()
        label = para.add_run("假动作剔除：")
        label.font.bold = True
        for action in fake_actions:
            bullet = doc.add_paragraph(style='List Bullet')
            bullet.add_run(action)
    
    # 止损点
    stop_loss = optimize_data.get('stop_loss', [])
    if stop_loss:
        para = doc.add_paragraph()
        label = para.add_run("止损点设置：")
        label.font.bold = True
        for item in stop_loss:
            bullet = doc.add_paragraph(style='List Bullet')
            bullet.add_run(item)
    
    doc.add_paragraph()


def create_checklist_section(doc):
    """创建快速思考清单"""
    add_heading_custom(doc, "快速思考清单", level=2)
    
    checklist = [
        "1. 定目标：这件事的终局是什么？（Why & What）",
        "2. 拆动作：第一步能做什么？（Actionable）",
        "3. 排优先级：哪件事价值最大？（80/20）",
        "4. 看结果：做完这一步离目标更近了吗？（Feedback）"
    ]
    
    for item in checklist:
        para = doc.add_paragraph(item)
        para.paragraph_format.line_spacing = 1.5
    
    doc.add_paragraph()


def generate_report(data, output_path):
    """生成完整的效率提升方案"""
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.size = Pt(11)
    
    # 提取数据
    task_name = data.get('task_name', '')
    date_str = data.get('date', datetime.now().strftime('%Y年%m月%d日'))
    
    # 1. 封面
    create_cover(doc, task_name, date_str)
    
    # 2. 问题界定
    add_heading_custom(doc, "第一步：界定问题", level=1)
    doc.add_paragraph()
    
    clarify = data.get('clarify', {})
    if clarify:
        create_5w1h_table(doc, clarify)
        
        # 最坏结果检验
        worst_case = clarify.get('worst_case', {})
        if worst_case:
            add_heading_custom(doc, "最坏结果检验", level=3)
            for key, value in worst_case.items():
                para = doc.add_paragraph()
                label = para.add_run(f"{key}：")
                label.font.bold = True
                para.add_run(str(value))
    
    # 分页
    doc.add_page_break()
    
    # 3. 任务拆解
    add_heading_custom(doc, "第二步：拆解任务", level=1)
    doc.add_paragraph()
    
    deconstruct = data.get('deconstruct', {})
    if deconstruct:
        create_wbs_section(doc, deconstruct)
        
        # MECE检查
        mece = deconstruct.get('mece_check', {})
        if mece:
            add_heading_custom(doc, "MECE检查", level=3)
            for key, value in mece.items():
                para = doc.add_paragraph()
                label = para.add_run(f"{key}：")
                label.font.bold = True
                para.add_run(str(value))
    
    # 分页
    doc.add_page_break()
    
    # 4. 路径优化
    add_heading_custom(doc, "第三步：优化路径", level=1)
    doc.add_paragraph()
    
    optimize = data.get('optimize', {})
    if optimize:
        create_optimization_section(doc, optimize)
    
    # 分页
    doc.add_page_break()
    
    # 5. 动态复盘
    add_heading_custom(doc, "第四步：动态复盘", level=1)
    doc.add_paragraph()
    
    review = data.get('review', {})
    if review:
        # PDCA说明
        pdca = review.get('pdca', {})
        if pdca:
            add_heading_custom(doc, "PDCA循环", level=2)
            for key, value in pdca.items():
                para = doc.add_paragraph()
                label = para.add_run(f"{key}：")
                label.font.bold = True
                para.add_run(str(value))
        
        # 归因分析
        attribution = review.get('attribution', {})
        if attribution:
            add_heading_custom(doc, "归因分析", level=2)
            for key, value in attribution.items():
                para = doc.add_paragraph()
                label = para.add_run(f"{key}：")
                label.font.bold = True
                para.add_run(str(value))
    
    # 分页
    doc.add_page_break()
    
    # 6. 执行手册
    add_heading_custom(doc, "执行手册", level=1)
    doc.add_paragraph()
    
    create_checklist_section(doc)
    
    # 工具模板
    add_heading_custom(doc, "常用工具模板", level=2)
    
    templates = [
        ("5W1H提问模板", "用于界定问题"),
        ("WBS分解模板", "用于拆解任务"),
        ("二八分析模板", "用于优化路径"),
        ("PDCA复盘模板", "用于动态复盘"),
        ("SOP固化模板", "用于经验沉淀")
    ]
    
    for name, desc in templates:
        para = doc.add_paragraph(style='List Bullet')
        label = para.add_run(f"{name}")
        label.font.bold = True
        para.add_run(f" - {desc}")
    
    # 保存文档
    doc.save(output_path)
    print(f"[OK] Efficiency boost plan generated: {output_path}")


def main():
    """主函数"""
    if len(sys.argv) < 2:
        print("Usage: python generate_report.py <json_file> [output_path]")
        print("Example: python generate_report.py data.json efficiency_plan.docx")
        sys.exit(1)
    
    json_file = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else "效率提升方案.docx"
    
    # 读取JSON数据
    with open(json_file, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    # 生成报告
    generate_report(data, output_path)


if __name__ == '__main__':
    main()
