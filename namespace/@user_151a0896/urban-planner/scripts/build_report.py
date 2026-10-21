#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_report.py —— 城市规划成果生成器

读取一个结构化 JSON，同时生成 Word(docx) 与 PPT(pptx) 规划成果文件；
若 JSON 含 `regulatory` 字段，另导出 GIS 可导入的控规指标 CSV；
若 JSON 含 `four_lines` 字段，自动生成「四线管控」章节；
运行时会依据 indicators.md / regulatory-plan.md 的阈值对 `regulatory` 地块做
「指标越界 / 千人指标配套」合规性自动核对，并导出 `_合规核对.csv`。

多场景适配：通过 --scenario 切换 PPT 强调色与封面标注
（government 政府汇报 / review 专家评审 / investment 招商推介）。

用法:
    python build_report.py report.json [--outdir DIR]
                                    [--no-docx] [--no-pptx] [--no-csv]
                                    [--body-font 字体] [--head-font 字体]
                                    [--scenario standard|government|review|investment]

JSON 结构见 assets/outline-template.md 或 references/agent-workflow.md。
依赖: pip install -r requirements.txt（位于 scripts/ 目录；仅做 docx/pptx 可不装场地分析依赖）
"""
import os
import sys
import json
import csv
import argparse
import platform

# 运行时不写 __pycache__，避免生成 .pyc 被打包/上传工具误判为二进制
sys.dont_write_bytecode = True

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from pptx import Presentation
    from pptx.util import Inches as PInches, Pt as PPt
    from pptx.oxml.ns import qn as Pqn
    from pptx.dml.color import RGBColor as PRGBColor
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


def _install_hint():
    # 给出可直接复制的安装命令，降低外部库安装门槛
    req = os.path.join(os.path.dirname(os.path.abspath(__file__)), "requirements.txt")
    if os.path.exists(req):
        return f'pip install -r "{req}"'
    return "pip install python-docx python-pptx"

# 场景 → 强调色(RGB 元组) / 标签（多场景适配：政府汇报/专家评审/招商推介）
SCENARIOS = {
    'standard':   {'accent': (0x1F, 0x4E, 0x79), 'label': '通用版'},
    'government': {'accent': (0x1F, 0x4E, 0x79), 'label': '政府汇报版'},
    'review':     {'accent': (0x2E, 0x6B, 0x5E), 'label': '专家评审版'},
    'investment': {'accent': (0xC0, 0x50, 0x4D), 'label': '招商推介版'},
}
PACCENT = None        # pptx 专用强调色（pptx 可用时于 main 内赋值）
PACCENT_DX = None     # docx 专用强调色（docx 可用时于 main 内赋值）
SCENARIO_LABEL = SCENARIOS['standard']['label']

# ---------- 字体（按操作系统自动选择，可用命令行参数覆盖） ----------
def _default_body_font():
    s = platform.system()
    if s == "Windows":
        return "宋体"
    if s == "Darwin":
        return "Songti SC"
    return "Noto Sans CJK SC"


def _default_head_font():
    s = platform.system()
    if s == "Windows":
        return "微软雅黑"
    if s == "Darwin":
        return "PingFang SC"
    return "Noto Sans CJK SC"


BODY_FONT = _default_body_font()
HEAD_FONT = _default_head_font()

MAX_PER_SLIDE = 9       # PPT 每页正文要点上限（超出自动分页）

# 控规指标表列（同时用于 docx 表头与 CSV 表头）
REG_COLS = ['地块编号', '用地代码', '用地性质', '用地面积(㎡)', '容积率',
            '建筑密度(%)', '建筑高度(m)', '绿地率(%)', '配套设施', '备注']


# ---------- 数值解析 ----------
def _to_float(v):
    try:
        return float(str(v).replace('%', '').strip())
    except (ValueError, TypeError):
        return None


# ---------- 合规性自动核对（指标越界 / 千人指标配套） ----------
# 阈值来源：references/indicators.md、references/regulatory-plan.md（策划参考值）
def check_compliance(data):
    """轻量合规提醒：绿地率底线（新区≥30%/旧区≥25%）+ 居住教育配套千人指标
    + 容积率/建筑密度的合理性 sanity 校验（非强制性，越界/异常给提示）。"""
    warns = []
    for r in data.get('regulatory') or []:
        code = str(r.get('use_code') or '') + str(r.get('use_name') or '')
        gid = r.get('code', '?')
        green = _to_float(r.get('green'))
        if green is not None:
            if green < 25:
                warns.append(f"[{gid}] 绿地率 {green}% 低于旧区底线 25%，越界！")
            elif green < 30:
                warns.append(f"[{gid}] 绿地率 {green}% 低于新区底线 30%，建议校核。")
        density = _to_float(r.get('density'))
        if density is not None and density > 60:
            warns.append(f"[{gid}] 建筑密度 {density}% 异常偏高（常规 ≤35%~50%），请校核。")
        far = _to_float(r.get('far'))
        if far is not None and far > 8:
            warns.append(f"[{gid}] 容积率 {far} 异常偏高，请校核（常规居住 ≤5、商业酌情）。")
        facility = str(r.get('facility') or '')
        if code.startswith('R') or '居住' in code:
            if not any(k in facility for k in ['幼儿园', '小学', '中学', '九年一贯', '教育', '学校', '托']):
                warns.append(f"[{gid}] 居住用地未标注教育配套，请按千人指标校核"
                             f"（幼儿园 40/千人、小学 80/千人、中学 80/千人）。")
    return warns


# ---------- 场地简报注入（site_analyze.py 产物） ----------
def inject_site(data, brief):
    """把 site_analyze.py 生成的 site_brief.json 合并进规划成果。"""
    merge = brief.get('merge') or {}
    for b in merge.get('basis') or []:
        data.setdefault('basis', []).append(b)
    for ind in merge.get('indicators') or []:
        data.setdefault('indicators', []).append(ind)
    loc = brief.get('location') or {}
    geo = brief.get('geometry') or {}
    climate = brief.get('climate')
    sug = brief.get('suggestions') or []
    surroundings = brief.get('surroundings') or {}
    site_assess = brief.get('site_assessment') or []
    blocks = [
        {'type': 'p', 'text': f"地块位于{loc.get('province','')}{loc.get('city','')}（{loc.get('address','') or '—'}），坐标约 {loc.get('lat')}, {loc.get('lng')}。"},
        {'type': 'p', 'text': f"红线面积约 {geo.get('area_m2')}㎡，周长约 {geo.get('perimeter_m')}m，主轴朝向约 {geo.get('orientation_deg')}°（正北为 0°）。"},
        {'type': 'p', 'text': f"气候区：{climate or '—'}。"},
    ]
    if site_assess:
        blocks.append({'type': 'h', 'text': '场地研判'})
        blocks.append({'type': 'bullets', 'items': site_assess})
    if surroundings.get('ok'):
        blocks.append({'type': 'h', 'text': '周边配套（15 分钟生活圈，OpenStreetMap）'})
        blocks.append({'type': 'bullets', 'items': surroundings.get('summary', '').splitlines()})
    elif surroundings.get('summary'):
        blocks.append({'type': 'p', 'text': surroundings.get('summary')})
    if sug:
        blocks.append({'type': 'h', 'text': '规划建议（来自场地分析）'})
        blocks.append({'type': 'bullets', 'items': sug})
    blocks.append({'type': 'p', 'text': '交互地图见同目录 site_map.html，可在浏览器打开查看红线落位、定位标记与周边路网（OSM 底图）。'})
    data.setdefault('sections', []).insert(0, {'chapter': '场地定位与地图', 'blocks': blocks})


# ---------- CJK 字体设置 ----------
def set_cjk(run, font_name):
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), font_name)
    rfonts.set(qn('w:ascii'), font_name)
    rfonts.set(qn('w:hAnsi'), font_name)


def set_cjk_ppt(run, font_name):
    run.font.name = font_name
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(Pqn('a:ea'))
    if ea is None:
        ea = rpr.makeelement(Pqn('a:ea'), {})
        rpr.append(ea)
    ea.set('typeface', font_name)
    latin = rpr.find(Pqn('a:latin'))
    if latin is None:
        latin = rpr.makeelement(Pqn('a:latin'), {})
        rpr.append(latin)
    latin.set('typeface', font_name)


# ================= DOCX =================
def build_docx(data, out_path):
    doc = Document()
    # 默认正文字体
    normal = doc.styles['Normal']
    normal.font.name = BODY_FONT
    normal.font.size = Pt(12)
    n_rpr = normal.element.get_or_add_rPr()
    n_fonts = n_rpr.find(qn('w:rFonts'))
    if n_fonts is None:
        n_fonts = n_rpr.makeelement(qn('w:rFonts'), {})
        n_rpr.append(n_fonts)
    n_fonts.set(qn('w:eastAsia'), BODY_FONT)
    n_fonts.set(qn('w:ascii'), BODY_FONT)
    n_fonts.set(qn('w:hAnsi'), BODY_FONT)

    title = data.get('title', '规划报告')
    subtitle = data.get('subtitle', '')
    client = data.get('client', '')
    date = data.get('date', '')

    # 封面
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(text=title)
    r.bold = True
    r.font.size = Pt(22)
    set_cjk(r, HEAD_FONT)
    r.font.color.rgb = PACCENT_DX
    if subtitle:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rs = s.add_run(text=subtitle)
        rs.font.size = Pt(15)
        set_cjk(rs, HEAD_FONT)
    for line, size in [(client, 12), (date, 12)]:
        if line:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = p.add_run(text=line)
            rr.font.size = Pt(size)
            set_cjk(rr, BODY_FONT)
    if SCENARIO_LABEL and SCENARIO_LABEL != '通用版':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(text=SCENARIO_LABEL)
        rr.font.size = Pt(12)
        rr.bold = True
        rr.font.color.rgb = PACCENT_DX
        set_cjk(rr, HEAD_FONT)
    doc.add_page_break()

    # 规划依据
    basis = data.get('basis') or []
    if basis:
        h = doc.add_heading('规划依据', level=1)
        for rn in h.runs:
            set_cjk(rn, HEAD_FONT)
        for b in basis:
            p = doc.add_paragraph(style='List Bullet')
            r = p.add_run(text=b)
            set_cjk(r, BODY_FONT)

    # 章节
    for sec in data.get('sections', []):
        ch = doc.add_heading(sec.get('chapter', ''), level=1)
        for rn in ch.runs:
            set_cjk(rn, HEAD_FONT)
        intro = sec.get('intro')
        if intro:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(24)
            r = p.add_run(text=intro)
            set_cjk(r, BODY_FONT)
        for blk in sec.get('blocks', []):
            bt = blk.get('type')
            if bt == 'h':
                hh = doc.add_heading(blk.get('text', ''), level=2)
                for rn in hh.runs:
                    set_cjk(rn, HEAD_FONT)
            elif bt == 'p':
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(24)
                p.paragraph_format.line_spacing = 1.5
                r = p.add_run(text=blk.get('text', ''))
                set_cjk(r, BODY_FONT)
            elif bt == 'bullets':
                for it in blk.get('items', []):
                    p = doc.add_paragraph(style='List Bullet')
                    r = p.add_run(text=it)
                    set_cjk(r, BODY_FONT)
            elif bt == 'numbered':
                for it in blk.get('items', []):
                    p = doc.add_paragraph(style='List Number')
                    r = p.add_run(text=it)
                    set_cjk(r, BODY_FONT)

    # 指标体系表
    inds = data.get('indicators') or []
    if inds:
        h = doc.add_heading('指标体系', level=1)
        for rn in h.runs:
            set_cjk(rn, HEAD_FONT)
        cols = ['类别', '指标', '单位', '近期', '远期', '属性']
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = 'Light Grid Accent 1'
        hdr = table.rows[0].cells
        for i, c in enumerate(cols):
            hdr[i].text = ''
            run = hdr[i].paragraphs[0].add_run(text=c)
            run.bold = True
            set_cjk(run, HEAD_FONT)
        for row in inds:
            cells = table.add_row().cells
            vals = [row.get('category', ''), row.get('name', ''), row.get('unit', ''),
                    row.get('recent', ''), row.get('future', ''), row.get('attr', '')]
            for i, v in enumerate(vals):
                cells[i].text = ''
                run = cells[i].paragraphs[0].add_run(text=str(v))
                set_cjk(run, BODY_FONT)

    # 重点项目库表
    projs = data.get('projects') or []
    if projs:
        h = doc.add_heading('重点项目库', level=1)
        for rn in h.runs:
            set_cjk(rn, HEAD_FONT)
        cols = ['项目名称', '建设内容', '规模', '时序', '责任主体']
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = 'Light Grid Accent 1'
        hdr = table.rows[0].cells
        for i, c in enumerate(cols):
            hdr[i].text = ''
            run = hdr[i].paragraphs[0].add_run(text=c)
            run.bold = True
            set_cjk(run, HEAD_FONT)
        for row in projs:
            cells = table.add_row().cells
            vals = [row.get('name', ''), row.get('content', ''), row.get('scale', ''),
                    row.get('period', ''), row.get('owner', '')]
            for i, v in enumerate(vals):
                cells[i].text = ''
                run = cells[i].paragraphs[0].add_run(text=str(v))
                set_cjk(run, BODY_FONT)

    # 控制性详细规划指标表
    regs = data.get('regulatory') or []
    if regs:
        h = doc.add_heading('控制性详细规划指标表', level=1)
        for rn in h.runs:
            set_cjk(rn, HEAD_FONT)
        table = doc.add_table(rows=1, cols=len(REG_COLS))
        table.style = 'Light Grid Accent 1'
        hdr = table.rows[0].cells
        for i, c in enumerate(REG_COLS):
            hdr[i].text = ''
            run = hdr[i].paragraphs[0].add_run(text=c)
            run.bold = True
            set_cjk(run, HEAD_FONT)
        for row in regs:
            cells = table.add_row().cells
            vals = [row.get('code', ''), row.get('use_code', ''), row.get('use_name', ''),
                    row.get('area', ''), row.get('far', ''), row.get('density', ''),
                    row.get('height', ''), row.get('green', ''), row.get('facility', ''),
                    row.get('note', '')]
            for i, v in enumerate(vals):
                cells[i].text = ''
                run = cells[i].paragraphs[0].add_run(text=str(v))
                set_cjk(run, BODY_FONT)

    # 四线管控章节（自动生成并标注管控要求）
    fls = data.get('four_lines') or []
    if fls:
        h = doc.add_heading('四线管控', level=1)
        for rn in h.runs:
            set_cjk(rn, HEAD_FONT)
        cols = ['管控线', '管控范围', '管控要求']
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = 'Light Grid Accent 1'
        hdr = table.rows[0].cells
        for i, c in enumerate(cols):
            hdr[i].text = ''
            run = hdr[i].paragraphs[0].add_run(text=c)
            run.bold = True
            set_cjk(run, HEAD_FONT)
        for row in fls:
            cells = table.add_row().cells
            vals = [row.get('type', ''), row.get('scope', ''), row.get('requirement', '')]
            for i, v in enumerate(vals):
                cells[i].text = ''
                run = cells[i].paragraphs[0].add_run(text=str(v))
                set_cjk(run, BODY_FONT)

    doc.save(out_path)


# ================= PPTX =================
def _add_one_slide(prs, title, items, size_head=28, size_body=16):
    """新增一页纯文本幻灯片。items: list of (text, is_head)。"""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(PInches(0.8), PInches(0.5), PInches(11.7), PInches(6.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = PPt(size_head)
    r.bold = True
    r.font.color.rgb = PACCENT
    set_cjk_ppt(r, HEAD_FONT)
    for txt, is_head in items:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = txt
        r.font.size = PPt(18 if is_head else size_body)
        if is_head:
            r.bold = True
        set_cjk_ppt(r, BODY_FONT)


def _add_paginated(prs, title, items, per=MAX_PER_SLIDE):
    """新增一页或多页（内容超出 per 时自动分页，续页标题带「（续 N）」）。"""
    if not items:
        _add_one_slide(prs, title, [])
        return
    chunks = [items[i:i + per] for i in range(0, len(items), per)]
    for idx, chunk in enumerate(chunks):
        t = title if (len(chunks) == 1 or idx == 0) else f"{title}（续 {idx}）"
        _add_one_slide(prs, t, chunk)


def build_pptx(data, out_path):
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    blank = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[0]

    title = data.get('title', '规划报告')
    subtitle = data.get('subtitle', '')
    client = data.get('client', '')
    date = data.get('date', '')

    # 封面
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = title
    set_cjk_ppt(s.shapes.title.text_frame.paragraphs[0].runs[0], HEAD_FONT)
    if s.placeholders and len(s.placeholders) > 1:
        sub = s.placeholders[1]
        sub.text = '\n'.join([x for x in [subtitle,
                                          SCENARIO_LABEL if SCENARIO_LABEL != '通用版' else '',
                                          client, date] if x])
        for para in sub.text_frame.paragraphs:
            for run in para.runs:
                set_cjk_ppt(run, BODY_FONT)

    # 目录
    chapters = [sec.get('chapter', '') for sec in data.get('sections', [])]
    if chapters:
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(PInches(0.8), PInches(0.5), PInches(11.7), PInches(6.5))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = '目录'
        r.font.size = PPt(32)
        r.bold = True
        r.font.color.rgb = PACCENT
        set_cjk_ppt(r, HEAD_FONT)
        for i, ch in enumerate(chapters, 1):
            p = tf.add_paragraph()
            r = p.add_run()
            r.text = f'{i}. {ch}'
            r.font.size = PPt(20)
            set_cjk_ppt(r, BODY_FONT)

    # 每章一页（必要时分页）
    for sec in data.get('sections', []):
        bullets = []
        for blk in sec.get('blocks', []):
            bt = blk.get('type')
            if bt == 'h':
                bullets.append(('• ' + blk.get('text', ''), True))
            elif bt == 'p':
                txt = blk.get('text', '')
                bullets.append((txt[:60] + ('…' if len(txt) > 60 else ''), False))
            elif bt == 'bullets':
                for it in blk.get('items', [])[:4]:
                    bullets.append((it, False))
            elif bt == 'numbered':
                for it in blk.get('items', [])[:4]:
                    bullets.append((it, False))
        _add_paginated(prs, sec.get('chapter', ''), bullets)

    # 指标体系页（分页）
    inds = data.get('indicators') or []
    if inds:
        items = []
        for row in inds:
            txt = f"· {row.get('name','')}（{row.get('unit','')}）：近期 {row.get('recent','')} / 远期 {row.get('future','')}（{row.get('attr','')}）"
            items.append((txt, False))
        _add_paginated(prs, '指标体系', items, per=12)

    # 重点项目页（分页）
    projs = data.get('projects') or []
    if projs:
        items = []
        for row in projs:
            txt = f"· {row.get('name','')}：{row.get('content','')}（{row.get('period','')}）"
            items.append((txt, False))
        _add_paginated(prs, '重点项目库', items, per=12)

    # 控规指标表页（分页）
    regs = data.get('regulatory') or []
    if regs:
        items = []
        for r in regs:
            items.append((
                f"· {r.get('code','')}｜{r.get('use_code','')} {r.get('use_name','')}"
                f"｜FAR {r.get('far','')} 密度{r.get('density','')}% 高{r.get('height','')}m 绿{r.get('green','')}%",
                False))
        _add_paginated(prs, '控规指标表', items, per=12)

    # 四线管控页
    fls = data.get('four_lines') or []
    if fls:
        items = []
        for r in fls:
            items.append((f"· {r.get('type', '')}：{r.get('requirement', '')}", False))
        _add_paginated(prs, '四线管控', items, per=12)

    prs.save(out_path)


# ================= CSV（控规属性表，GIS 可导入） =================
def build_csv(data, out_path):
    regs = data.get('regulatory') or []
    if not regs:
        return False
    # utf-8-sig：Excel 直接打开中文不乱码
    with open(out_path, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow(REG_COLS)
        for r in regs:
            w.writerow([
                r.get('code', ''), r.get('use_code', ''), r.get('use_name', ''),
                r.get('area', ''), r.get('far', ''), r.get('density', ''),
                r.get('height', ''), r.get('green', ''), r.get('facility', ''),
                r.get('note', ''),
            ])
    return True


# ================= 入口 =================
def main():
    ap = argparse.ArgumentParser(description='生成城市规划 docx/pptx/csv 成果')
    ap.add_argument('json', help='输入 JSON 文件路径')
    ap.add_argument('--outdir', default='.', help='输出目录')
    ap.add_argument('--no-docx', action='store_true', help='不生成 docx')
    ap.add_argument('--no-pptx', action='store_true', help='不生成 pptx')
    ap.add_argument('--no-csv', action='store_true', help='即使含 regulatory 也不导出 CSV')
    ap.add_argument('--body-font', default=None,
                    help='正文字体名（默认按系统自动选择：Windows 宋体 / macOS Songti SC / Linux Noto Sans CJK SC）')
    ap.add_argument('--head-font', default=None,
                    help='标题字体名（默认按系统自动选择：Windows 微软雅黑 / macOS PingFang SC / Linux Noto Sans CJK SC）')
    ap.add_argument('--scenario', default='standard',
                    choices=['standard', 'government', 'review', 'investment'],
                    help='成果场景：standard 通用 / government 政府汇报 / review 专家评审 / investment 招商推介（影响 PPT 配色与封面标注）')
    ap.add_argument('--site', default=None,
                    help='注入 site_analyze.py 生成的 site_brief.json：自动合并属地政策到 basis、规划底线到 indicators，并生成「场地定位与地图」章节')
    args = ap.parse_args()

    if not DOCX_AVAILABLE and not args.no_docx:
        print(f'[warn] python-docx 未安装，跳过 docx 生成。请运行：{_install_hint()}')
    if not PPTX_AVAILABLE and not args.no_pptx:
        print(f'[warn] python-pptx 未安装，跳过 pptx 生成。请运行：{_install_hint()}')

    global BODY_FONT, HEAD_FONT, PACCENT, PACCENT_DX, SCENARIO_LABEL
    BODY_FONT = args.body_font or _default_body_font()
    HEAD_FONT = args.head_font or _default_head_font()
    _sc = SCENARIOS.get(args.scenario, SCENARIOS['standard'])
    if PPTX_AVAILABLE:
        PACCENT = PRGBColor(*_sc['accent'])
    if DOCX_AVAILABLE:
        PACCENT_DX = RGBColor(*_sc['accent'])
    SCENARIO_LABEL = _sc['label']

    try:
        with open(args.json, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except FileNotFoundError:
        sys.exit(f'[错误] 找不到输入文件：{args.json}')
    except json.JSONDecodeError as e:
        sys.exit(f'[错误] JSON 解析失败（{args.json}）：{e}')

    # 注入场地简报（site_analyze.py 产物）
    if args.site:
        try:
            with open(args.site, 'r', encoding='utf-8') as f:
                brief = json.load(f)
            inject_site(data, brief)
            print(f'[场地] 已注入 site_brief：{args.site}')
        except Exception as e:
            print(f'[警告] 注入场地简报失败（{args.site}）：{e}')

    os.makedirs(args.outdir, exist_ok=True)
    safe = ''.join(c for c in data.get('title', '规划成果') if c not in r'\\/:*?"<>|')
    if not safe:
        safe = '规划成果'

    if not args.no_docx and DOCX_AVAILABLE:
        dx = os.path.join(args.outdir, safe + '.docx')
        build_docx(data, dx)
        print('已生成:', dx)
    if not args.no_pptx and PPTX_AVAILABLE:
        px = os.path.join(args.outdir, safe + '.pptx')
        build_pptx(data, px)
        print('已生成:', px)
    if not args.no_csv and build_csv(data, os.path.join(args.outdir, safe + '_控规指标表.csv')):
        print('已生成:', os.path.join(args.outdir, safe + '_控规指标表.csv'))

    # 合规性自动核对（指标越界 / 千人指标配套）
    warns = check_compliance(data)
    if warns:
        print('\n[合规核对] 指标越界 / 千人指标配套 提示：')
        for w in warns:
            print('  -', w)
        if not args.no_csv:
            cpath = os.path.join(args.outdir, safe + '_合规核对.csv')
            with open(cpath, 'w', encoding='utf-8-sig', newline='') as f:
                wcsv = csv.writer(f)
                wcsv.writerow(['提示项'])
                for w in warns:
                    wcsv.writerow([w])
            print('已生成:', cpath)
    else:
        print('\n[合规核对] 控规指标未触发越界/配套预警。')


if __name__ == '__main__':
    main()
